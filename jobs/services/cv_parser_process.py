"""Small, resource-bounded parser process. Receives bytes; emits only plain text."""
import io
import json
import re
import sys
import zipfile

MAX_INPUT = 5 * 1024 * 1024
MAX_CHARS = 30000


def parse(data, name):
    if name.lower().endswith('.pdf') or data[:5] == b'%PDF-':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = '\n'.join((page.extract_text() or '') for page in reader.pages[:12])
    elif name.lower().endswith('.docx') or data[:2] == b'PK':
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 500 or sum(e.file_size for e in entries) > 20 * 1024 * 1024:
                return ''
        import docx
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables[:20]:
            for row in table.rows[:500]:
                parts.append(' | '.join(cell.text for cell in row.cells))
        text = '\n'.join(parts)
    else:
        return ''
    return re.sub(r'\n{3,}', '\n\n', re.sub(r'[ \t]+', ' ', text)).strip()[:MAX_CHARS]


if __name__ == '__main__':
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (8, 8))
        if sys.platform.startswith('linux'):
            resource.setrlimit(resource.RLIMIT_AS, (384 * 1024 * 1024, 384 * 1024 * 1024))
        data = sys.stdin.buffer.read(MAX_INPUT + 1)
        text = parse(data, sys.argv[1]) if len(data) <= MAX_INPUT else ''
    except Exception:
        text = ''
    sys.stdout.write(json.dumps(text))
